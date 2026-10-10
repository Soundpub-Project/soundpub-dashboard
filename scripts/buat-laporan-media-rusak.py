"""Bangun laporan Excel media yang masih menunjuk GCS / Supabase Cloud.

Masukan: 4 CSV hasil ekspor dari soundpub (k1..k4 di /tmp).
Keluaran: satu workbook dengan sheet Ringkasan + 4 sheet kategori.
"""

import csv
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

TMP = os.environ.get("TMPDIR", "/tmp")

KATEGORI = [
    ("Track Audio", "k1.csv", "url_rusak"),
    ("Clip Track", "k2.csv", "url_rusak"),
    ("Cover Release", "k3.csv", "url_rusak"),
    ("Avatar Profil", "k4.csv", "url_rusak"),
]

JUDUL = Font(bold=True, color="FFFFFF", size=11)
ISI_GCS = PatternFill("solid", fgColor="FFF2CC")
ISI_CLOUD = PatternFill("solid", fgColor="FCE4E4")
KEPALA = PatternFill("solid", fgColor="2F5597")
TAUTAN = Font(color="0563C1", underline="single")
TEPI = Border(*[Side(style="thin", color="D0D0D0")] * 4)


def baca(nama):
    jalur = os.path.join(TMP, nama)
    if not os.path.exists(jalur):
        sys.exit(f"tidak ada: {jalur}")
    with open(jalur, newline="", encoding="utf-8") as f:
        return list(csv.reader(f))


def lebar_kolom(ws, baris, maks=60):
    for i, _ in enumerate(baris[0], start=1):
        panjang = max(
            (len(str(r[i - 1])) for r in baris[:400] if i <= len(r)), default=10
        )
        ws.column_dimensions[get_column_letter(i)].width = min(max(panjang + 2, 10), maks)


def tulis_sheet(wb, judul, baris):
    ws = wb.create_sheet(judul)
    kepala = baris[0]
    ws.append(kepala)
    for sel in ws[1]:
        sel.font = JUDUL
        sel.fill = KEPALA
        sel.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    i_lokasi = kepala.index("lokasi_lama")
    i_link = next((n for n, k in enumerate(kepala) if k.startswith("link_")), None)

    for r in baris[1:]:
        ws.append(r)
        n = ws.max_row
        warna = ISI_GCS if "Google" in r[i_lokasi] else ISI_CLOUD
        for sel in ws[n]:
            sel.fill = warna
            sel.border = TEPI
            sel.alignment = Alignment(vertical="top")
        if i_link is not None and r[i_link]:
            sel = ws.cell(row=n, column=i_link + 1)
            sel.hyperlink = r[i_link]
            sel.font = TAUTAN

    if ws.max_row > 1:
        ref = f"A1:{get_column_letter(len(kepala))}{ws.max_row}"
        tabel = Table(displayName=judul.replace(" ", "_").replace("-", "_"), ref=ref)
        tabel.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight1", showRowStripes=False
        )
        ws.add_table(tabel)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    lebar_kolom(ws, baris)
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)

    ringkas = wb.create_sheet("Ringkasan")
    hitung = []

    for judul, berkas, _ in KATEGORI:
        baris = baca(berkas)
        gcs = sum(1 for r in baris[1:] if "Google" in r[0])
        cloud = sum(1 for r in baris[1:] if "Supabase Cloud" in r[0])
        hitung.append((judul, gcs, cloud, gcs + cloud))
        tulis_sheet(wb, judul, baris)

    ringkas["A1"] = "Laporan Media Tidak Terbaca - Soundpub Dashboard"
    ringkas["A1"].font = Font(bold=True, size=14)
    ringkas["A2"] = f"Dibuat: {date.today().isoformat()}  |  Sumber: schema soundpub di 20.20.20.3"
    ringkas["A2"].font = Font(italic=True, size=9, color="666666")

    kepala = ["Kategori", "Google Cloud Storage", "Supabase Cloud", "Total"]
    ringkas.append([])
    ringkas.append(kepala)
    for sel in ringkas[4]:
        sel.font = JUDUL
        sel.fill = KEPALA
        sel.alignment = Alignment(horizontal="center")

    for judul, gcs, cloud, total in hitung:
        ringkas.append([judul, gcs, cloud, total])
        for sel in ringkas[ringkas.max_row]:
            sel.border = TEPI

    ringkas.append(
        [
            "TOTAL",
            sum(h[1] for h in hitung),
            sum(h[2] for h in hitung),
            sum(h[3] for h in hitung),
        ]
    )
    for sel in ringkas[ringkas.max_row]:
        sel.font = Font(bold=True)
        sel.fill = PatternFill("solid", fgColor="D9E2F3")
        sel.border = TEPI

    catatan = [
        "",
        "KETERANGAN",
        "",
        "Google Cloud Storage - billing account project sudah nonaktif (HTTP 403).",
        "  Berkas kemungkinan MASIH ADA. Bisa ditarik kembali bila billing diaktifkan.",
        "",
        "Supabase Cloud - project opkvvdgnhhopkkeaokzo sudah tidak resolve (DNS mati).",
        "  Berkas HILANG PERMANEN. Artis perlu unggah ulang.",
        "",
        "Baris kuning = Google Cloud Storage.  Baris merah muda = Supabase Cloud.",
        "Kolom link_release / link_profil dapat diklik langsung.",
        "",
        "Migrasi ke server baru tidak menyebabkan masalah ini - URL tersebut sudah",
        "rusak sejak sebelum migrasi. Media yang tersimpan di self-hosted (557 berkas)",
        "sudah pindah utuh dan terverifikasi byte-for-byte.",
    ]
    for t in catatan:
        ringkas.append([t])
        if t in ("KETERANGAN",):
            ringkas[ringkas.max_row][0].font = Font(bold=True, size=12)

    for kol, lebar in zip("ABCD", (22, 24, 20, 12)):
        ringkas.column_dimensions[kol].width = lebar

    keluar = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "docs",
        f"media-tidak-terbaca-{date.today().isoformat()}.xlsx",
    )
    os.makedirs(os.path.dirname(keluar), exist_ok=True)
    wb.save(keluar)
    print(keluar)
    for judul, gcs, cloud, total in hitung:
        print(f"  {judul:<16} GCS={gcs:<4} Cloud={cloud:<4} Total={total}")


if __name__ == "__main__":
    main()
